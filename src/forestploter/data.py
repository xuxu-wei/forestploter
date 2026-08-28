"""Long-table data objects and XLSX/CSV readers for forest plots.

The reader keeps file-format details outside the renderer. XLSX input retains
cell values, record order, and supported vertical merges. CSV input uses the
same logical fields but cannot carry merge information. Excel fonts, colors,
row heights, and column widths are not imported as plot styles.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd


@dataclass(frozen=True)
class ForestDataMapping:
    """Map the nine reserved long-table fields to source column names.

    Parameters
    ----------
    plot_row : str, default "_plot_row"
        Visual-row identifier. Records with the same value use exactly the
        same vertical coordinate.
    series : str, default "_series"
        Series identifier used for style lookup and legend entries.
    ci_column : str, default "_ci_column"
        Target CI-column identifier. It must match the ``key`` of a
        :class:`forestploter.ForestColumn` whose ``role`` is ``"ci"``.
    estimate : str, default "estimate"
        Point-estimate field.
    lower : str, default "lower"
        Confidence-interval lower-bound field.
    upper : str, default "upper"
        Confidence-interval upper-bound field.
    row_type : str, default "_row_type"
        Row-semantic field. Accepted values are ``"header"``,
        ``"estimate"``, ``"summary"``, and ``"spacer"``.
    indent : str, default "_indent"
        Non-negative indentation-level field.
    is_summary : str, default "_is_summary"
        Strict boolean field that controls whether the observation uses the
        summary marker.

    Notes
    -----
    Prefer the default names. A mapping only renames fields; it does not change
    the long-table semantics. In particular, it cannot map several wide-form
    triples to several series. One source record always contains at most one
    ``estimate/lower/upper`` triple.
    """

    plot_row: str = "_plot_row"
    series: str = "_series"
    ci_column: str = "_ci_column"
    estimate: str = "estimate"
    lower: str = "lower"
    upper: str = "upper"
    row_type: str = "_row_type"
    indent: str = "_indent"
    is_summary: str = "_is_summary"

    def technical_fields(self) -> tuple[str, ...]:
        """Return reserved fields in the recommended physical-file order.

        Returns
        -------
        tuple of str
            The three statistical fields followed by the six control fields.
            User-defined display fields are not included.
        """

        return (
            self.estimate,
            self.lower,
            self.upper,
            self.plot_row,
            self.series,
            self.ci_column,
            self.row_type,
            self.indent,
            self.is_summary,
        )


@dataclass(frozen=True)
class ForestCellSpan:
    """Describe one supported vertical merge in an XLSX display field.

    Parameters
    ----------
    column_key : str
        Display-field name containing the merge.
    start_record : int
        Zero-based index of the first source record in the merge.
    end_record : int
        Zero-based inclusive index of the final source record in the merge.

    Notes
    -----
    Rendering also verifies that each merge covers complete ``_plot_row``
    groups. Reserved fields, headers, horizontal merges, and two-dimensional
    merges are unsupported.
    """

    column_key: str
    start_record: int
    end_record: int


@dataclass(frozen=True)
class ForestData:
    """Store a renderable long table and its source-cell structure.

    Parameters
    ----------
    frame : pandas.DataFrame
        Source long-table values. Each estimate record describes one series in
        one CI column with one statistical triple.
    spans : tuple of ForestCellSpan, default ()
        Supported XLSX vertical merges. CSV input always produces an empty
        tuple.
    mapping : ForestDataMapping, default ForestDataMapping()
        Reserved-field mapping.
    source : pathlib.Path or None, default None
        Absolute input path. It may be ``None`` when constructed directly from
        a DataFrame.
    sheet_name : str or None, default None
        XLSX worksheet name; ``None`` for CSV.
    source_rows : tuple of int, default ()
        One-based source-file row number for every record.
    cell_references : tuple of tuple of str, default ()
        Source references with the same shape as ``frame``, for example
        ``"Forest!D7"`` or ``"D7"`` for CSV.

    Notes
    -----
    XLSX “what you see is what you enter” applies to values, record order, and
    supported vertical merges. The plot does not copy Excel fonts, fills,
    borders, column widths, or row heights. Give each series its own
    ``_plot_row``. To show one series in several CI columns, duplicate its
    record, keep the same ``_plot_row``, and change ``_ci_column``; all CI
    columns then share the same vertical coordinate.
    """

    frame: pd.DataFrame
    spans: tuple[ForestCellSpan, ...] = ()
    mapping: ForestDataMapping = field(default_factory=ForestDataMapping)
    source: Path | None = None
    sheet_name: str | None = None
    source_rows: tuple[int, ...] = ()
    cell_references: tuple[tuple[str, ...], ...] = ()

    def cell_reference(self, record: int, column: str) -> str | None:
        """Return the source-cell reference for one record and field.

        Parameters
        ----------
        record : int
            Zero-based record index.
        column : str
            Field name in ``frame``.

        Returns
        -------
        str or None
            Source-cell reference, or ``None`` when no source references were
            supplied.

        Raises
        ------
        IndexError
            Raised when ``record`` is outside the data range.
        KeyError
            Raised when ``column`` does not exist.
        """

        if record < 0 or record >= len(self.frame):
            raise IndexError("Forest record index is out of range.")
        try:
            column_index = list(self.frame.columns).index(column)
        except ValueError as error:
            raise KeyError(column) from error
        if not self.cell_references:
            return None
        return self.cell_references[record][column_index]


def _validate_mapping(mapping: ForestDataMapping) -> None:
    names = mapping.technical_fields()
    if any(not isinstance(name, str) or not name.strip() for name in names):
        raise ValueError("Forest data mapping fields must be non-empty strings.")
    if len(names) != len(set(names)):
        raise ValueError("Forest data mapping fields must be unique.")


def _trim_trailing_empty_rows(values: list[list[Any]]) -> list[list[Any]]:
    while values and all(pd.isna(value) or value == "" for value in values[-1]):
        values.pop()
    return values


def _read_csv(
    source: Path,
    *,
    header_row: int,
    mapping: ForestDataMapping,
    encoding: str,
) -> ForestData:
    frame = pd.read_csv(
        source,
        header=header_row - 1,
        encoding=encoding,
        skip_blank_lines=False,
    )
    while not frame.empty and frame.iloc[-1].isna().all():
        frame = frame.iloc[:-1]
    frame = frame.reset_index(drop=True)
    if frame.empty:
        raise ValueError("Forest data must contain at least one record.")
    blank_rows = frame.index[frame.isna().all(axis=1)].tolist()
    if blank_rows:
        row_number = header_row + 1 + blank_rows[0]
        raise ValueError(
            f"CSV contains an entirely blank data row at line {row_number}; "
            "use an explicit row_type='spacer' record instead."
        )
    if any(str(name).startswith("Unnamed:") for name in frame.columns):
        raise ValueError("Forest CSV header contains an empty field name.")
    if len(frame.columns) != len(set(frame.columns)):
        raise ValueError("Forest data field names must be unique.")
    rows = tuple(header_row + 1 + index for index in range(len(frame)))
    refs = tuple(
        tuple(
            f"{_column_letter(column_index + 1)}{row}" for column_index in range(len(frame.columns))
        )
        for row in rows
    )
    return ForestData(
        frame=frame,
        mapping=mapping,
        source=source.resolve(),
        source_rows=rows,
        cell_references=refs,
    )


def _column_letter(index: int) -> str:
    letters = ""
    value = index
    while value:
        value, remainder = divmod(value - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _select_worksheet(workbook: Any, sheet_name: str | int) -> Any:
    if isinstance(sheet_name, bool):
        raise TypeError("sheet_name must be a worksheet name or zero-based integer index.")
    if isinstance(sheet_name, int):
        if sheet_name < 0 or sheet_name >= len(workbook.worksheets):
            raise ValueError(f"XLSX sheet index is out of range: {sheet_name}")
        return workbook.worksheets[sheet_name]
    if isinstance(sheet_name, str):
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"XLSX worksheet does not exist: {sheet_name!r}")
        return workbook[sheet_name]
    raise TypeError("sheet_name must be a worksheet name or zero-based integer index.")


def _read_xlsx(
    source: Path,
    *,
    sheet_name: str | int,
    header_row: int,
    mapping: ForestDataMapping,
) -> ForestData:
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - dependency failure is environment-specific
        raise ImportError(
            "Reading XLSX files requires the optional dependency 'openpyxl'."
        ) from error

    cached_book = load_workbook(source, data_only=True, read_only=False)
    formula_book = load_workbook(source, data_only=False, read_only=False)
    try:
        cached_sheet = _select_worksheet(cached_book, sheet_name)
        formula_sheet = formula_book[cached_sheet.title]
        header_values = [
            cached_sheet.cell(header_row, column).value
            for column in range(1, cached_sheet.max_column + 1)
        ]
        while header_values and (header_values[-1] is None or str(header_values[-1]).strip() == ""):
            header_values.pop()
        if not header_values:
            raise ValueError(f"XLSX header row {header_row} is empty.")
        if any(value is None or not str(value).strip() for value in header_values):
            raise ValueError("XLSX header row contains an empty field name.")
        headers = [str(value).strip() for value in header_values]
        if len(headers) != len(set(headers)):
            raise ValueError("Forest data field names must be unique.")

        raw_values = [
            [cached_sheet.cell(row, column).value for column in range(1, len(headers) + 1)]
            for row in range(header_row + 1, cached_sheet.max_row + 1)
        ]
        raw_values = _trim_trailing_empty_rows(raw_values)
        if not raw_values:
            raise ValueError("Forest data must contain at least one record.")
        for offset, values in enumerate(raw_values):
            if all(value is None or value == "" for value in values):
                row_number = header_row + 1 + offset
                raise ValueError(
                    f"XLSX contains an entirely blank data row at {cached_sheet.title}!{row_number}; "
                    "use an explicit row_type='spacer' record instead."
                )

        technical = set(mapping.technical_fields())
        for row_offset, values in enumerate(raw_values):
            row_number = header_row + 1 + row_offset
            for column_index, header in enumerate(headers, start=1):
                formula_cell = formula_sheet.cell(row_number, column_index)
                if formula_cell.data_type == "f" and header in technical:
                    cached_value = values[column_index - 1]
                    if cached_value is None or cached_value == "":
                        raise ValueError(
                            f"Required formula at {cached_sheet.title}!{formula_cell.coordinate} has no cached result; "
                            "recalculate and save the workbook in Excel before reading it."
                        )

        spans: list[ForestCellSpan] = []
        header_lookup = {index + 1: name for index, name in enumerate(headers)}
        for merged in cached_sheet.merged_cells.ranges:
            if merged.max_row <= header_row:
                continue
            if merged.min_row <= header_row:
                raise ValueError(f"Merged range {merged} crosses or includes the XLSX header row.")
            if merged.min_col != merged.max_col or merged.min_row == merged.max_row:
                raise ValueError(
                    f"Merged range {merged} is unsupported; only single-column vertical merges below the header are allowed."
                )
            if merged.min_col > len(headers):
                raise ValueError(f"Merged range {merged} lies outside the forest data fields.")
            column_key = header_lookup[merged.min_col]
            if column_key in technical:
                raise ValueError(
                    f"Merged range {merged} targets technical field {column_key!r}; "
                    "only display fields may be merged."
                )
            start = merged.min_row - header_row - 1
            end = merged.max_row - header_row - 1
            if start < 0 or end >= len(raw_values):
                raise ValueError(f"Merged range {merged} extends outside the forest data records.")
            spans.append(ForestCellSpan(column_key, start, end))

        frame = pd.DataFrame(raw_values, columns=headers)
        rows = tuple(header_row + 1 + index for index in range(len(frame)))
        refs = tuple(
            tuple(
                f"{cached_sheet.title}!{_column_letter(column_index + 1)}{row}"
                for column_index in range(len(headers))
            )
            for row in rows
        )
        return ForestData(
            frame=frame,
            spans=tuple(
                sorted(
                    spans, key=lambda item: (item.column_key, item.start_record, item.end_record)
                )
            ),
            mapping=mapping,
            source=source.resolve(),
            sheet_name=cached_sheet.title,
            source_rows=rows,
            cell_references=refs,
        )
    finally:
        cached_book.close()
        formula_book.close()


def read_forest_data(
    source: str | Path,
    *,
    sheet_name: str | int = 0,
    header_row: int = 1,
    mapping: ForestDataMapping | None = None,
    encoding: str = "utf-8-sig",
) -> ForestData:
    """Read an XLSX or CSV file that follows the forest long-table contract.

    Parameters
    ----------
    source : str or pathlib.Path
        Path to an ``.xlsx`` or ``.csv`` file. Other suffixes are not guessed.
    sheet_name : str or int, default 0
        XLSX worksheet name or zero-based index. CSV has no worksheets and
        rejects any value other than the default ``0``.
    header_row : int, default 1
        One-based physical row containing field names. Every non-empty physical
        row below it is one source record.
    mapping : ForestDataMapping, optional
        Reserved-field name mapping. The nine standard names are used by
        default.
    encoding : str, default "utf-8-sig"
        CSV text encoding. Ignored for XLSX.

    Returns
    -------
    ForestData
        Long-table values, supported XLSX merges, and traceable source-cell
        references.

    Raises
    ------
    FileNotFoundError
        Raised when ``source`` does not exist.
    TypeError
        Raised for invalid ``header_row``, ``sheet_name``, or ``mapping``
        types.
    ValueError
        Raised when the file type, header, worksheet, merge, or cached formula
        result violates the contract.
    ImportError
        Raised when XLSX support is requested without ``openpyxl``.

    Notes
    -----
    Use a first worksheet named ``Forest`` and one row of unique field names.
    Two orders are deliberately independent:

    * The ``columns`` argument of :func:`forestploter.forest` is the only thing
      that determines the final left-to-right plot order. Display columns may
      appear before or after any CI column.
    * The XLSX/CSV header order exists for readable data entry and does not
      determine plot order. A recommended file places display fields that
      precede the first CI first, then ``estimate/lower/upper`` as the physical
      representation of that CI, then trailing display fields, and finally the
      six controls ``_plot_row``, ``_series``, ``_ci_column``, ``_row_type``,
      ``_indent``, and ``_is_summary``.

    Enter only one statistical triple per estimate record. For three series,
    use three records and normally three different ``_plot_row`` values. A
    shared outcome cell may be vertically merged across those records, while
    an unmerged effect-text field can contain one value per series; every text
    row is then exactly aligned with its CI. To show the same series in two CI
    columns, duplicate its record, use a different ``_ci_column``, and retain
    the same ``_plot_row``. The two plotted intervals will share one y position.

    Merges are accepted only in display fields and must be single-column,
    vertical ranges below the header. Merges in reserved fields, horizontal or
    rectangular merges, and merges crossing the header are rejected. Blank
    values are never forward-filled. All three of ``estimate/lower/upper`` may
    be blank to skip an interval; a partially blank triple fails validation in
    :func:`forestploter.forest`.

    The package does not evaluate Excel formulas. It uses cached results saved
    in the workbook and asks the user to recalculate and save when a required
    formula has no cache. CSV follows the same record and ``_plot_row`` rules,
    but cannot represent merges. Excel fonts, fills, borders, row heights, and
    column widths do not override :class:`forestploter.ForestTheme` or
    :class:`forestploter.ForestColumn` settings.

    Examples
    --------
    >>> from forestploter import read_forest_data
    >>> df = read_forest_data("forest_input.xlsx", sheet_name="Forest")  # doctest: +SKIP
    >>> df.frame.columns[:5].tolist()  # doctest: +SKIP
    ['Outcome', 'estimate', 'lower', 'upper', 'Effect [95% CI]']
    """

    path = Path(source)
    if not path.exists():
        raise FileNotFoundError(path)
    if isinstance(header_row, bool) or not isinstance(header_row, int):
        raise TypeError("header_row must be a positive one-based integer.")
    if header_row < 1:
        raise ValueError("header_row must be a positive one-based integer.")
    resolved_mapping = mapping or ForestDataMapping()
    if not isinstance(resolved_mapping, ForestDataMapping):
        raise TypeError("mapping must be a ForestDataMapping instance.")
    _validate_mapping(resolved_mapping)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        if sheet_name != 0:
            raise ValueError(
                "CSV input does not support sheet_name; leave it at the default value 0."
            )
        return _read_csv(
            path,
            header_row=header_row,
            mapping=resolved_mapping,
            encoding=encoding,
        )
    if suffix == ".xlsx":
        return _read_xlsx(
            path,
            sheet_name=sheet_name,
            header_row=header_row,
            mapping=resolved_mapping,
        )
    raise ValueError("Forest data source must be an .xlsx or .csv file.")


__all__ = ["ForestCellSpan", "ForestData", "ForestDataMapping", "read_forest_data"]
