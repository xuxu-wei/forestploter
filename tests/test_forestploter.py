"""forestploter 的文件契约、布局和视觉回归测试。"""

from __future__ import annotations

import inspect
from dataclasses import replace
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

import forestploter
from forestploter import (
    ForestCellSpan,
    ForestColumn,
    ForestData,
    ForestDataMapping,
    ForestLegendSpec,
    ForestReferenceLegendSpec,
    ForestSeriesStyle,
    ForestTheme,
    compute_column_geometry,
    forest,
    read_forest_data,
)
from forestploter.gallery_cases import GALLERY_FUNCTIONS
from tests.visual_cases import (
    CASE_BUILDERS,
    CASE_CSV_FILES,
    CASE_DATA_FILES,
    DATA_DIR,
    artifact_name_for_case,
    build_case,
    load_case_data,
)

TEST_OUTPUT_DIR = Path(__file__).resolve().parent / ".cache" / "pytest-artifacts"

STATISTIC_FIELDS = ("estimate", "lower", "upper")
CONTROL_FIELDS = (
    "_plot_row",
    "_series",
    "_ci_column",
    "_row_type",
    "_indent",
    "_is_summary",
)
EXPECTED_FILE_HEADERS = {
    "single_series": ("label", "n", *STATISTIC_FIELDS, "effect_display", *CONTROL_FIELDS),
    "multi_series": ("outcome", "participants", *STATISTIC_FIELDS, *CONTROL_FIELDS),
    "dual_ci_columns": ("endpoint", "n", *STATISTIC_FIELDS, *CONTROL_FIELDS),
    "clipping_stress": (
        "scenario",
        "expected_clipping",
        "interval_display",
        *STATISTIC_FIELDS,
        *CONTROL_FIELDS,
    ),
    "long_layout": (
        "subgroup",
        "region",
        "participants",
        "effect_display",
        *STATISTIC_FIELDS,
        *CONTROL_FIELDS,
    ),
    "auto_scale_mixed_effects": (
        "outcome",
        "scale",
        "n",
        "effect_display",
        *STATISTIC_FIELDS,
        *CONTROL_FIELDS,
    ),
    "four_series_dense": ("endpoint", "participants", *STATISTIC_FIELDS, *CONTROL_FIELDS),
    "two_by_two_ci_columns": (
        "outcome",
        "n",
        *STATISTIC_FIELDS,
        "cohort",
        *CONTROL_FIELDS,
    ),
    "deep_hierarchy_many_rows": (
        "label",
        "n",
        "effect_display",
        *STATISTIC_FIELDS,
        *CONTROL_FIELDS,
    ),
    "unicode_custom_theme": (
        "结局",
        "人群",
        "样本量",
        "效应值",
        *STATISTIC_FIELDS,
        *CONTROL_FIELDS,
    ),
    "boundary_precision": (
        "scenario",
        "position",
        "interval_display",
        *STATISTIC_FIELDS,
        *CONTROL_FIELDS,
    ),
    "multi_series_ci_text_rows": (
        "outcome",
        "participants",
        *STATISTIC_FIELDS,
        "series_ci_text",
        *CONTROL_FIELDS,
    ),
    "comprehensive_showcase": (
        "label",
        "participants",
        *STATISTIC_FIELDS,
        "crude_ci_text",
        "adjusted_ci_text",
        "note",
        *CONTROL_FIELDS,
    ),
}
EXPECTED_MERGE_COUNTS = {
    "single_series": 0,
    "multi_series": 16,
    "dual_ci_columns": 10,
    "clipping_stress": 0,
    "long_layout": 0,
    "auto_scale_mixed_effects": 0,
    "four_series_dense": 14,
    "two_by_two_ci_columns": 24,
    "deep_hierarchy_many_rows": 0,
    "unicode_custom_theme": 0,
    "boundary_precision": 0,
    "multi_series_ci_text_rows": 14,
    "comprehensive_showcase": 0,
}


@pytest.mark.parametrize("case_name", tuple(CASE_BUILDERS))
def test_xlsx_visual_case_renders_with_valid_layout(case_name: str) -> None:
    data = load_case_data(case_name)
    result = build_case(case_name, data)
    try:
        result.figure.canvas.draw()
        diagnostics = result.layout_diagnostics
        assert len(result.row_centers) == len(diagnostics.plot_row_ids)
        assert len(diagnostics.row_bounds) == len(diagnostics.plot_row_ids)
        assert len(diagnostics.row_scales) == len(diagnostics.plot_row_ids)
        assert diagnostics.header_overflow_count == 0
        assert diagnostics.series_gap_used == 0.0
        assert diagnostics.final_figure_width > 0
        assert result.geometry[0].left == pytest.approx(0.0)
        assert result.geometry[-1].right == pytest.approx(1.0)
        assert all(item.left < item.right for item in result.geometry)
        assert len(diagnostics.observation_positions) == sum(data.frame["estimate"].notna())
    finally:
        plt.close(result.figure)


@pytest.mark.parametrize("case_name", tuple(CASE_BUILDERS))
def test_csv_companion_has_equivalent_rendering_semantics(case_name: str) -> None:
    xlsx_data = load_case_data(case_name, source_format="xlsx")
    csv_data = load_case_data(case_name, source_format="csv")
    xlsx_result = build_case(case_name, xlsx_data)
    csv_result = build_case(case_name, csv_data)
    try:
        assert csv_data.spans == ()
        assert (
            xlsx_result.layout_diagnostics.plot_row_ids
            == csv_result.layout_diagnostics.plot_row_ids
        )
        xlsx_positions = xlsx_result.layout_diagnostics.observation_positions
        csv_positions = csv_result.layout_diagnostics.observation_positions
        assert [item[:3] for item in xlsx_positions] == [item[:3] for item in csv_positions]
        assert [item[3] for item in xlsx_positions] == pytest.approx(
            [item[3] for item in csv_positions]
        )
        assert xlsx_result.row_centers == pytest.approx(csv_result.row_centers)
        assert xlsx_result.clipped_intervals == csv_result.clipped_intervals
    finally:
        plt.close(xlsx_result.figure)
        plt.close(csv_result.figure)


def test_column_geometry_preserves_relative_widths() -> None:
    geometry = compute_column_geometry(
        (
            ForestColumn("label", "Label", "text", 2.0),
            ForestColumn("ci", "Effect", "ci", 3.0),
        )
    )
    assert geometry[0].left == pytest.approx(0.0)
    assert geometry[0].right == pytest.approx(0.4)
    assert geometry[1].left == pytest.approx(0.4)
    assert geometry[1].right == pytest.approx(1.0)


def test_mapping_reports_statistics_before_the_six_controls() -> None:
    assert ForestDataMapping().technical_fields() == (*STATISTIC_FIELDS, *CONTROL_FIELDS)


def test_data_csv_and_image_names_have_a_one_to_one_mapping() -> None:
    assert CASE_DATA_FILES.keys() == CASE_BUILDERS.keys() == CASE_CSV_FILES.keys()
    image_names = [artifact_name_for_case(case_name) for case_name in CASE_BUILDERS]
    assert len(image_names) == len(set(image_names))
    for case_name, xlsx_name in CASE_DATA_FILES.items():
        stem = Path(xlsx_name).stem
        assert CASE_CSV_FILES[case_name] == f"{stem}.csv"
        assert artifact_name_for_case(case_name) == f"{stem}.png"
        assert (DATA_DIR / xlsx_name).exists()
        assert (DATA_DIR / CASE_CSV_FILES[case_name]).exists()


@pytest.mark.parametrize("case_name", tuple(CASE_BUILDERS))
def test_example_headers_follow_readable_file_order(case_name: str) -> None:
    """CSV and XLSX use the same explicit entry order with controls last."""

    xlsx = load_case_data(case_name, source_format="xlsx")
    csv = load_case_data(case_name, source_format="csv")
    expected = EXPECTED_FILE_HEADERS[case_name]
    assert tuple(xlsx.frame.columns) == expected
    assert tuple(csv.frame.columns) == expected
    assert tuple(xlsx.frame.columns[-len(CONTROL_FIELDS) :]) == CONTROL_FIELDS
    assert tuple(csv.frame.columns[-len(CONTROL_FIELDS) :]) == CONTROL_FIELDS


@pytest.mark.parametrize("case_name", tuple(CASE_BUILDERS))
def test_example_xlsx_types_merges_and_validated_fields(case_name: str) -> None:
    """Every workbook preserves types, expected merges, and entry validation."""

    path = DATA_DIR / CASE_DATA_FILES[case_name]
    data = read_forest_data(path)
    assert len(data.spans) == EXPECTED_MERGE_COUNTS[case_name]
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = workbook["Forest"]
        header_index = {
            str(cell.value): index
            for index, cell in enumerate(sheet[1], start=1)
            if cell.value is not None
        }
        for row in range(2, sheet.max_row + 1):
            for field in STATISTIC_FIELDS:
                value = sheet.cell(row, header_index[field]).value
                assert value is None or (
                    isinstance(value, (int, float)) and not isinstance(value, bool)
                )
            indent = sheet.cell(row, header_index["_indent"]).value
            summary = sheet.cell(row, header_index["_is_summary"]).value
            assert isinstance(indent, (int, float)) and not isinstance(indent, bool)
            assert isinstance(summary, bool)
        validations = list(sheet.data_validations.dataValidation)
        formulas = " ".join(str(item.formula1) for item in validations)
        assert all(value in formulas for value in ("header", "estimate", "summary", "spacer"))
        assert "TRUE" in formulas and "FALSE" in formulas
    finally:
        workbook.close()


def test_gallery_functions_are_the_shared_df_entry_points() -> None:
    assert GALLERY_FUNCTIONS is CASE_BUILDERS
    assert tuple(GALLERY_FUNCTIONS) == tuple(CASE_DATA_FILES)
    for function in GALLERY_FUNCTIONS.values():
        parameters = tuple(inspect.signature(function).parameters)
        assert parameters == ("df",)
        assert "read_forest_data" not in inspect.getsource(function)


def test_xlsx_reader_preserves_merge_spans_and_source_coordinates() -> None:
    data = load_case_data("multi_series_ci_text_rows")
    assert data.source is not None and data.source.suffix == ".xlsx"
    assert data.sheet_name == "Forest"
    assert len(data.spans) == 14
    assert {span.column_key for span in data.spans} == {"outcome", "participants"}
    assert data.cell_reference(0, "outcome") == "Forest!A2"
    first_outcome = next(span for span in data.spans if span.column_key == "outcome")
    assert (first_outcome.start_record, first_outcome.end_record) == (0, 2)
    assert data.frame.loc[0, "outcome"] == "Overall survival"
    assert data.frame.isna().loc[1, "outcome"]


def test_xlsx_reader_supports_named_sheet_and_nondefault_header_row(tmp_path: Path) -> None:
    path = tmp_path / "header-row.xlsx"
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme["A1"] = "Instructions"
    sheet = workbook.create_sheet("Forest")
    sheet["A1"] = "Forest example"
    sheet["A2"] = "Header follows on row 3"
    sheet.append(
        [
            "label",
            "_plot_row",
            "_series",
            "_ci_column",
            "estimate",
            "lower",
            "upper",
            "_row_type",
            "_indent",
            "_is_summary",
        ]
    )
    sheet.append(["Outcome A", "a", "Series A", "ci", 0.8, 0.6, 1.0, "estimate", 0, False])
    workbook.save(path)
    workbook.close()

    data = read_forest_data(path, sheet_name="Forest", header_row=3)
    assert data.sheet_name == "Forest"
    assert data.source_rows == (4,)
    assert data.cell_reference(0, "estimate") == "Forest!E4"
    assert data.frame.loc[0, "_plot_row"] == "a"


def test_template_has_forest_first_and_readme_second() -> None:
    template = DATA_DIR / "forest_data_template.xlsx"
    workbook = load_workbook(template, read_only=False, data_only=True)
    try:
        assert workbook.sheetnames == ["Forest", "README"]
        assert workbook["Forest"]["A1"].value == "Outcome"
        assert tuple(cell.value for cell in workbook["Forest"][1]) == (
            "Outcome",
            "Series",
            *STATISTIC_FIELDS,
            "Crude effect",
            "Adjusted effect",
            *CONTROL_FIELDS,
        )
        readme_text = "\n".join(
            str(cell.value)
            for row in workbook["README"].iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert "long-table template" in readme_text and "长表模板" in readme_text
        assert "Accepted values / 允许值" in readme_text
        assert "Plot order / 绘图顺序" in readme_text
        assert {str(item) for item in workbook["Forest"].merged_cells.ranges} == {
            "A2:A5",
            "B2:B3",
            "B4:B5",
        }
        validations = list(workbook["Forest"].data_validations.dataValidation)
        formulas = " ".join(str(item.formula1) for item in validations)
        assert all(value in formulas for value in ("header", "estimate", "summary", "spacer"))
        assert "TRUE" in formulas and "FALSE" in formulas
    finally:
        workbook.close()


def test_multi_series_each_text_record_shares_y_with_its_ci() -> None:
    data = load_case_data("multi_series_ci_text_rows")
    result = build_case("multi_series_ci_text_rows", data)
    try:
        assert tuple(item.key for item in result.geometry) == (
            "outcome",
            "participants",
            "ci",
            "series_ci_text",
        )
        by_plot_row = {
            plot_row: y
            for plot_row, _series, _ci, y in result.layout_diagnostics.observation_positions
        }
        text_artists = {
            artist.get_gid().removeprefix("cell-text:series_ci_text:"): artist
            for artist in result.axes.texts
            if artist.get_gid() and artist.get_gid().startswith("cell-text:series_ci_text:")
        }
        assert len(text_artists) == 21
        for plot_row, artist in text_artists.items():
            assert artist.get_position()[1] == pytest.approx(by_plot_row[plot_row])
            assert "[" in artist.get_text() and "]" in artist.get_text()
    finally:
        plt.close(result.figure)


def test_single_series_renders_text_ci_text_in_declared_order() -> None:
    result = build_case("single_series")
    try:
        assert tuple(item.key for item in result.geometry) == (
            "label",
            "n",
            "ci",
            "effect_display",
        )
        ci = next(item for item in result.geometry if item.key == "ci")
        artists = [
            artist
            for artist in result.axes.texts
            if (artist.get_gid() or "").startswith("cell-text:effect_display:")
        ]
        assert artists and all(artist.get_position()[0] > ci.right for artist in artists)
    finally:
        plt.close(result.figure)


def test_two_by_two_case_uses_24_records_12_rows_and_cross_ci_alignment() -> None:
    data = load_case_data("two_by_two_ci_columns")
    result = build_case("two_by_two_ci_columns", data)
    try:
        assert tuple(item.key for item in result.geometry) == (
            "outcome",
            "n",
            "ci_crude",
            "ci_adjusted",
            "cohort",
        )
        assert len(data.frame) == 24
        assert len(result.row_centers) == 12
        grouped: dict[str, list[tuple[str, str, float]]] = {}
        for plot_row, series, ci_column, y in result.layout_diagnostics.observation_positions:
            grouped.setdefault(plot_row, []).append((series, ci_column, y))
        assert len(grouped) == 12
        for observations in grouped.values():
            assert len(observations) == 2
            assert {item[1] for item in observations} == {"ci_crude", "ci_adjusted"}
            assert len({item[0] for item in observations}) == 1
            assert observations[0][2] == pytest.approx(observations[1][2])
    finally:
        plt.close(result.figure)


def test_comprehensive_showcase_combines_hierarchy_alignment_clipping_and_legends() -> None:
    data = load_case_data("comprehensive_showcase")
    result = build_case("comprehensive_showcase", data)
    try:
        assert tuple(item.key for item in result.geometry) == (
            "label",
            "participants",
            "ci_crude",
            "crude_ci_text",
            "ci_adjusted",
            "adjusted_ci_text",
            "note",
        )
        assert len(data.frame) == 35
        assert len(result.row_centers) == 20
        assert result.figure.get_size_inches()[1] < 8.0
        assert set(data.frame["_series"].dropna()) == {
            "Integrated care",
            "Digital support",
            "Usual care",
        }

        header_records = data.frame.loc[data.frame["_row_type"] == "header"]
        child_records = data.frame.loc[data.frame["_row_type"] != "header"]
        assert len(header_records) == 5
        assert header_records["_plot_row"].nunique() == 5
        assert header_records[["estimate", "lower", "upper"]].isna().all().all()
        assert header_records[["_series", "_ci_column"]].isna().all().all()
        assert header_records["_indent"].eq(0).all()
        assert header_records["_is_summary"].eq(False).all()
        assert header_records[["label", "participants", "note"]].notna().all().all()
        assert len(child_records) == 30
        assert child_records["_plot_row"].nunique() == 15
        assert child_records["_indent"].eq(1).all()
        assert child_records[["participants", "note"]].isna().all().all()
        assert data.frame["_row_type"].value_counts().to_dict() == {
            "estimate": 24,
            "summary": 6,
            "header": 5,
        }
        visual_rows = data.frame.drop_duplicates("_plot_row", keep="first")
        assert visual_rows.groupby("_row_type", sort=False).size().to_dict() == {
            "header": 5,
            "estimate": 12,
            "summary": 3,
        }
        assert visual_rows["_row_type"].tolist() == (
            ["header", "estimate", "estimate", "estimate"] * 4
            + ["header", "summary", "summary", "summary"]
        )
        assert data.frame.groupby("_plot_row", sort=False).size().tolist() == [1, 2, 2, 2] * 5
        assert child_records.loc[child_records["_ci_column"] == "ci_crude", "label"].notna().all()
        assert child_records.loc[child_records["_ci_column"] == "ci_adjusted", "label"].isna().all()
        for field in ("crude_ci_text", "adjusted_ci_text"):
            ci_text = data.frame[field].dropna().astype(str)
            assert len(ci_text) == 15
            assert ci_text.str.match(r"^\d+\.\d{2} \[\d+\.\d{2}, \d+\.\d{2}\]$").all()
        assert (
            data.frame.loc[data.frame["crude_ci_text"].notna(), "_ci_column"].eq("ci_crude").all()
        )
        assert (
            data.frame.loc[data.frame["adjusted_ci_text"].notna(), "_ci_column"]
            .eq("ci_adjusted")
            .all()
        )

        grouped: dict[str, list[tuple[str, str, float]]] = {}
        for plot_row, series, ci_column, y in result.layout_diagnostics.observation_positions:
            grouped.setdefault(plot_row, []).append((series, ci_column, y))
        assert len(grouped) == 15
        for observations in grouped.values():
            assert len(observations) == 2
            assert {item[1] for item in observations} == {"ci_crude", "ci_adjusted"}
            assert len({item[0] for item in observations}) == 1
            assert observations[0][2] == pytest.approx(observations[1][2])

        expected_y = {observations[0][2] for observations in grouped.values()}
        for field in ("crude_ci_text", "adjusted_ci_text"):
            artists = [
                artist
                for artist in result.axes.texts
                if (artist.get_gid() or "").startswith(f"cell-text:{field}:")
            ]
            assert len(artists) == 15
            assert sorted(artist.get_position()[1] for artist in artists) == pytest.approx(
                sorted(expected_y)
            )

        assert data.spans == ()
        header_label = next(
            artist for artist in result.axes.texts if artist.get_gid() == "cell-text:label:c13-h001"
        )
        child_label = next(
            artist
            for artist in result.axes.texts
            if artist.get_gid() == "cell-text:label:c13-r001-integrated"
        )
        assert child_label.get_position()[0] > header_label.get_position()[0]
        assert result.clipped_intervals == 9
        gids = {
            artist.get_gid()
            for artist in [*result.axes.lines, *result.axes.patches, *result.axes.collections]
            if artist.get_gid()
        }
        assert not any(
            (artist.get_gid() or "").startswith("merged-") for artist in result.axes.texts
        )
        left_arrows = {gid for gid in gids if gid.endswith(":arrow:left")}
        right_arrows = {gid for gid in gids if gid.endswith(":arrow:right")}
        assert len(left_arrows) == 3
        assert len(right_arrows) == 6
        for arrow_gid in left_arrows | right_arrows:
            prefix, side = arrow_gid.rsplit(":arrow:", maxsplit=1)
            assert f"{prefix}:cap:{side}" not in gids

        assert {"reference-line:ci_crude", "reference-line:ci_adjusted"} <= gids
        assert "ideal-line:ci_crude" not in gids
        assert "ideal-line:ci_adjusted" in gids
        assert [item[0] for item in result.layout_diagnostics.legend_bounds] == [
            "series",
            "reference",
        ]
        assert all(item[4] < 0 for item in result.layout_diagnostics.legend_bounds)
        labels = {artist.get_text() for artist in result.axes.texts}
        assert {"Favours integrated care", "Favours usual care"} <= labels
    finally:
        plt.close(result.figure)


def test_ideal_line_columns_default_filter_and_validation() -> None:
    data = load_case_data("dual_ci_columns")
    columns = (
        ForestColumn("endpoint", "Endpoint", "text", 2.5),
        ForestColumn("n", "N", "numeric", 0.6, "right"),
        ForestColumn("ci_30d", "30-day", "ci", 3.0, "center"),
        ForestColumn("ci_90d", "90-day", "ci", 3.0, "center"),
    )
    default_result = forest(
        data,
        columns=columns,
        xlim=(0.4, 1.6),
        ref_line=1.0,
        ideal_line=0.75,
    )
    filtered_result = forest(
        data,
        columns=columns,
        xlim=(0.4, 1.6),
        ref_line=1.0,
        ideal_line=0.75,
        ideal_line_columns=("ci_90d",),
    )
    suppressed_result = forest(
        data,
        columns=columns,
        xlim=(0.4, 1.6),
        ref_line=1.0,
        ideal_line=0.75,
        ideal_line_columns=(),
    )
    try:
        default_gids = {line.get_gid() for line in default_result.axes.lines if line.get_gid()}
        filtered_gids = {line.get_gid() for line in filtered_result.axes.lines if line.get_gid()}
        suppressed_gids = {
            line.get_gid() for line in suppressed_result.axes.lines if line.get_gid()
        }
        assert {"reference-line:ci_30d", "reference-line:ci_90d"} <= default_gids
        assert {"ideal-line:ci_30d", "ideal-line:ci_90d"} <= default_gids
        assert {"reference-line:ci_30d", "reference-line:ci_90d"} <= filtered_gids
        assert "ideal-line:ci_30d" not in filtered_gids
        assert "ideal-line:ci_90d" in filtered_gids
        assert not {gid for gid in suppressed_gids if gid.startswith("ideal-line:")}
    finally:
        plt.close(default_result.figure)
        plt.close(filtered_result.figure)
        plt.close(suppressed_result.figure)

    with pytest.raises(ValueError, match="not a string"):
        forest(
            data,
            columns=columns,
            xlim=(0.4, 1.6),
            ideal_line=0.75,
            ideal_line_columns="ci_30d",
        )
    with pytest.raises(ValueError, match="role='ci'"):
        forest(
            data,
            columns=columns,
            xlim=(0.4, 1.6),
            ideal_line=0.75,
            ideal_line_columns=("endpoint",),
        )
    with pytest.raises(ValueError, match="role='ci'"):
        forest(
            data,
            columns=columns,
            xlim=(0.4, 1.6),
            ideal_line=0.75,
            ideal_line_columns=("unknown",),
        )
    with pytest.raises(ValueError, match="requires ideal_line"):
        forest(
            data,
            columns=columns,
            xlim=(0.4, 1.6),
            ideal_line_columns=("ci_30d",),
        )
    with pytest.raises(ValueError, match="duplicate"):
        forest(
            data,
            columns=columns,
            xlim=(0.4, 1.6),
            ideal_line=0.75,
            ideal_line_columns=("ci_30d", "ci_30d"),
        )
    with pytest.raises(ValueError, match="non-empty strings"):
        forest(
            data,
            columns=columns,
            xlim=(0.4, 1.6),
            ideal_line=0.75,
            ideal_line_columns=("",),
        )
    with pytest.raises(ValueError, match="non-empty strings"):
        forest(
            data,
            columns=columns,
            xlim=(0.4, 1.6),
            ideal_line=0.75,
            ideal_line_columns=(1,),  # type: ignore[arg-type]
        )

    inactive_data = data.frame.loc[data.frame["_ci_column"] == "ci_30d"].reset_index(drop=True)
    with pytest.raises(ValueError, match="active columns"):
        forest(
            inactive_data,
            columns=columns,
            xlim=(0.4, 1.6),
            ideal_line=0.75,
            ideal_line_columns=("ci_90d",),
        )
    with pytest.raises(ValueError, match="active ideal-line target"):
        forest(
            data,
            columns=columns,
            xlim=(0.4, 1.6),
            ideal_line=0.75,
            ideal_line_columns=(),
            reference_legend=ForestReferenceLegendSpec(ideal_label="Target"),
        )


def test_dual_ci_columns_share_one_y_per_endpoint() -> None:
    result = build_case("dual_ci_columns")
    try:
        grouped: dict[str, set[float]] = {}
        for plot_row, _series, _ci, y in result.layout_diagnostics.observation_positions:
            grouped.setdefault(plot_row, set()).add(y)
        assert len(result.row_centers) == 6
        assert all(len(y_values) == 1 for y_values in grouped.values())
    finally:
        plt.close(result.figure)


def test_bottom_legend_without_column_spans_all_active_ci_columns() -> None:
    data = load_case_data("two_by_two_ci_columns")
    columns = (
        ForestColumn("outcome", "Outcome", "text", 2.4),
        ForestColumn("n", "N", "numeric", 0.6, "right"),
        ForestColumn("cohort", "Cohort", "text", 1.1),
        ForestColumn("ci_crude", "Crude", "ci", 3.4, "center"),
        ForestColumn("ci_adjusted", "Adjusted", "ci", 3.4, "center"),
    )
    result = forest(
        data,
        columns=columns,
        ref_line=1.0,
        xlim=(0.5, 1.5),
        series_styles={
            "Cohort A": ForestSeriesStyle("Cohort A", "#1F4E79", marker="s"),
            "Cohort B": ForestSeriesStyle("Cohort B", "#D97706", marker="o"),
        },
        legend=ForestLegendSpec(location="bottom", ncol=2),
    )
    try:
        bounds = result.layout_diagnostics.legend_bounds[0]
        crude = next(item for item in result.geometry if item.key == "ci_crude")
        adjusted = next(item for item in result.geometry if item.key == "ci_adjusted")
        assert (bounds[1] + bounds[3]) / 2 == pytest.approx((crude.left + adjusted.right) / 2)
        assert bounds[4] < 0
    finally:
        plt.close(result.figure)


def test_merged_display_cell_is_centered_and_internal_dividers_are_suppressed() -> None:
    result = build_case("multi_series_ci_text_rows")
    try:
        first_span = next(
            item for item in result.layout_diagnostics.resolved_spans if item[0] == "outcome"
        )
        _, _start_id, _end_id, bottom, top = first_span
        merged_artist = next(
            artist for artist in result.axes.texts if artist.get_gid() == "merged-text:outcome:0:2"
        )
        assert merged_artist.get_position()[1] == pytest.approx((bottom + top) / 2)
        outcome_geometry = next(item for item in result.geometry if item.key == "outcome")
        internal_boundaries = {
            result.layout_diagnostics.row_bounds[0][0],
            result.layout_diagnostics.row_bounds[1][0],
        }
        for line in result.axes.lines:
            x = np.asarray(line.get_xdata(), dtype=float)
            y = np.asarray(line.get_ydata(), dtype=float)
            if x.size != 2 or y.size != 2 or not np.isclose(y[0], y[1]):
                continue
            if any(np.isclose(y[0], boundary) for boundary in internal_boundaries):
                assert not (x.min() < outcome_geometry.right and x.max() > outcome_geometry.left)
    finally:
        plt.close(result.figure)


def test_clipped_intervals_count_affected_intervals_once() -> None:
    result = build_case("clipping_stress")
    try:
        assert result.clipped_intervals == 5
    finally:
        plt.close(result.figure)


def test_arrow_replaces_cap_on_each_clipped_side() -> None:
    result = build_case("clipping_stress")
    try:
        gids = {
            artist.get_gid()
            for artist in [*result.axes.lines, *result.axes.patches]
            if artist.get_gid()
        }
        # record 2: left endpoint clipped; record 3: right endpoint clipped
        assert "ci:2:arrow:left" in gids
        assert "ci:2:cap:left" not in gids
        assert "ci:2:cap:right" in gids
        assert "ci:3:arrow:right" in gids
        assert "ci:3:cap:right" not in gids
        assert "ci:3:cap:left" in gids
        # record 4 spans both sides: line and arrows, no caps
        assert "ci:4:line" in gids
        assert "ci:4:arrow:left" in gids and "ci:4:arrow:right" in gids
        assert "ci:4:cap:left" not in gids and "ci:4:cap:right" not in gids
    finally:
        plt.close(result.figure)


def test_fully_offscale_interval_draws_only_nearest_arrow() -> None:
    result = build_case("clipping_stress")
    try:
        gids = {
            artist.get_gid()
            for artist in [*result.axes.lines, *result.axes.patches, *result.axes.collections]
            if artist.get_gid()
        }
        for record, side in ((5, "left"), (6, "right")):
            assert f"ci:{record}:arrow:{side}:offscale" in gids
            assert f"ci:{record}:line" not in gids
            assert f"ci:{record}:cap:left" not in gids
            assert f"ci:{record}:cap:right" not in gids
            assert f"ci:{record}:marker" not in gids
    finally:
        plt.close(result.figure)


def test_boundary_precision_counts_only_true_overflows() -> None:
    result = build_case("boundary_precision")
    try:
        assert result.clipped_intervals == 2
    finally:
        plt.close(result.figure)


def test_header_and_bottom_legends_have_diagnostic_bounds() -> None:
    header_result = build_case("two_by_two_ci_columns")
    bottom_result = build_case("four_series_dense")
    try:
        header_bounds = header_result.layout_diagnostics.legend_bounds
        bottom_bounds = bottom_result.layout_diagnostics.legend_bounds
        assert len(header_bounds) == len(bottom_bounds) == 1
        assert header_bounds[0][0] == bottom_bounds[0][0] == "series"
        assert header_bounds[0][2] > 0
        assert bottom_bounds[0][4] < 0
    finally:
        plt.close(header_result.figure)
        plt.close(bottom_result.figure)


def test_series_and_reference_legends_stack_in_requested_order() -> None:
    data = load_case_data("multi_series")
    columns = (
        ForestColumn("outcome", "Outcome", "text", 2.8),
        ForestColumn("participants", "N", "numeric", 0.8, "right"),
        ForestColumn("ci", "Effect", "ci", 4.8, "center"),
    )
    styles = {
        name: ForestSeriesStyle(name, color)
        for name, color in zip(
            ("Treatment A", "Treatment B", "Treatment C"),
            ("#1F4E79", "#D97706", "#5B8C5A"),
            strict=True,
        )
    }
    result = forest(
        data,
        columns=columns,
        xlim=(0.4, 1.6),
        ref_line=1.0,
        series_styles=styles,
        legend=ForestLegendSpec(location="bottom", ncol=3),
        reference_legend=ForestReferenceLegendSpec(
            reference_label="No effect", location="bottom", ncol=1
        ),
    )
    try:
        bounds = result.layout_diagnostics.legend_bounds
        assert [item[0] for item in bounds] == ["series", "reference"]
        assert bounds[0][2] > bounds[1][4]
    finally:
        plt.close(result.figure)


def test_long_layout_wraps_headers_and_grows_figure() -> None:
    result = build_case("long_layout")
    try:
        diagnostics = result.layout_diagnostics
        wrapped = dict(diagnostics.wrapped_headers)
        assert diagnostics.final_figure_width > 7.5
        assert diagnostics.final_figure_width <= 16.0
        assert any("\n" in header for header in wrapped.values())
        assert diagnostics.header_overflow_count == 0
    finally:
        plt.close(result.figure)


def test_deep_hierarchy_produces_a_tall_readable_figure() -> None:
    data = load_case_data("deep_hierarchy_many_rows")
    result = build_case("deep_hierarchy_many_rows", data)
    try:
        assert len(data.frame) >= 25
        assert data.frame["_indent"].max() == 2
        assert result.figure.get_size_inches()[1] > 9.0
        assert len(result.row_centers) == len(data.frame)
    finally:
        plt.close(result.figure)


def test_plot_result_saves_nonempty_png_and_svg() -> None:
    TEST_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    png = TEST_OUTPUT_DIR / artifact_name_for_case("single_series")
    svg = TEST_OUTPUT_DIR / "01_single_series.svg"
    result = build_case("single_series")
    try:
        saved_png = result.save(png, dpi=100)
        saved_svg = result.save(svg)
        assert saved_png == png.resolve() and saved_svg == svg.resolve()
        assert saved_png.stat().st_size > 10_000
        assert saved_svg.stat().st_size > 10_000
        with Image.open(saved_png) as image:
            assert image.format == "PNG"
            assert image.width >= 800
            assert image.height >= 300
        assert "<svg" in saved_svg.read_text(encoding="utf-8")[:500]
    finally:
        plt.close(result.figure)


def test_summary_flag_must_be_boolean() -> None:
    data = load_case_data("single_series", source_format="csv")
    frame = data.frame.copy()
    frame["_is_summary"] = frame["_is_summary"].map({True: "True", False: "False"})
    with pytest.raises(ValueError, match="summary.*boolean"):
        build_case("single_series", frame)


def test_all_blank_interval_is_explicitly_skipped() -> None:
    frame = load_case_data("single_series", source_format="csv").frame.copy()
    record = frame.index[frame["_row_type"] == "estimate"][0]
    frame.loc[record, ["estimate", "lower", "upper"]] = np.nan
    result = build_case("single_series", frame)
    try:
        assert len(result.layout_diagnostics.observation_positions) == 6
        assert f"ci:{record}:line" not in {
            line.get_gid() for line in result.axes.lines if line.get_gid()
        }
    finally:
        plt.close(result.figure)


@pytest.mark.parametrize(
    "column",
    (
        ForestColumn("label", "Label", "unknown", 1),
        ForestColumn("label", "Label", "text", 1, "middle"),
    ),
)
def test_column_role_and_alignment_are_validated(column: ForestColumn) -> None:
    with pytest.raises(ValueError, match="Unsupported forest column"):
        compute_column_geometry((column, ForestColumn("ci", "CI", "ci", 1)))


@pytest.mark.parametrize(
    ("theme", "match"),
    (
        (ForestTheme(show_vertical_grid="yes"), "show_vertical_grid"),  # type: ignore[arg-type]
        (ForestTheme(header_fill="not-a-color"), "invalid colors"),
    ),
)
def test_theme_values_are_validated(theme: ForestTheme, match: str) -> None:
    data = load_case_data("single_series")
    columns = (
        ForestColumn("label", "Label", "text", 2),
        ForestColumn("n", "N", "numeric", 1, "right"),
        ForestColumn("effect_display", "Effect", "numeric", 2, "right"),
        ForestColumn("ci", "CI", "ci", 3),
    )
    with pytest.raises(ValueError, match=match):
        forest(data, columns=columns, xlim=(-1, 1), theme=theme)


@pytest.mark.parametrize(
    ("mutation", "match"),
    (
        ("invalid_row_type", "Unsupported forest row types"),
        ("negative_indent", "indentation"),
        ("inverted_interval", "does not contain its estimate"),
        ("partial_interval", "all present or all blank"),
        ("missing_estimate", "lacks required fields"),
        ("unknown_ci", "does not name a column"),
    ),
)
def test_invalid_long_table_inputs_are_rejected(mutation: str, match: str) -> None:
    frame = load_case_data("single_series", source_format="csv").frame.copy()
    estimate_index = frame.index[frame["_row_type"] == "estimate"][0]
    if mutation == "invalid_row_type":
        frame.loc[estimate_index, "_row_type"] = "detail"
    elif mutation == "negative_indent":
        frame.loc[estimate_index, "_indent"] = -1
    elif mutation == "inverted_interval":
        frame.loc[estimate_index, "lower"] = frame.loc[estimate_index, "estimate"] + 0.1
    elif mutation == "partial_interval":
        frame.loc[estimate_index, "lower"] = np.nan
    elif mutation == "missing_estimate":
        frame = frame.drop(columns="estimate")
    elif mutation == "unknown_ci":
        frame.loc[estimate_index, "_ci_column"] = "missing_ci"
    with pytest.raises(ValueError, match=match):
        build_case("single_series", frame)


def test_plot_row_records_must_be_contiguous() -> None:
    frame = load_case_data("dual_ci_columns", source_format="csv").frame.copy()
    reordered = frame.iloc[[0, 1, 3, 2, *range(4, len(frame))]].reset_index(drop=True)
    with pytest.raises(ValueError, match="must be contiguous"):
        build_case("dual_ci_columns", reordered)


@pytest.mark.parametrize(
    ("field", "value", "match"),
    (
        ("_series", "Other series", "inconsistent series"),
        ("_row_type", "summary", "inconsistent row_type"),
        ("_indent", 2, "inconsistent indentation"),
        ("_is_summary", True, "inconsistent summary flags"),
        ("endpoint", "Conflicting label", "conflicting values"),
    ),
)
def test_records_sharing_plot_row_must_share_row_metadata(
    field: str, value: object, match: str
) -> None:
    frame = load_case_data("dual_ci_columns", source_format="csv").frame.copy()
    # Records 1 and 2 describe the same endpoint in two CI columns.
    frame.loc[2, field] = value
    with pytest.raises(ValueError, match=match):
        build_case("dual_ci_columns", frame)


def test_duplicate_plot_row_ci_target_is_rejected() -> None:
    frame = load_case_data("dual_ci_columns", source_format="csv").frame.copy()
    frame.loc[2, "_ci_column"] = "ci_30d"
    with pytest.raises(ValueError, match="more than one record"):
        build_case("dual_ci_columns", frame)


def test_merge_span_must_cover_complete_plot_row_groups() -> None:
    data = load_case_data("dual_ci_columns", source_format="csv")
    malformed = ForestData(
        frame=data.frame,
        spans=(ForestCellSpan("endpoint", 0, 1),),
    )
    with pytest.raises(ValueError, match="last record of a plot-row group"):
        build_case("dual_ci_columns", malformed)


def test_series_style_mapping_requires_exact_observed_keys() -> None:
    data = load_case_data("multi_series")
    columns = (
        ForestColumn("outcome", "Outcome", "text", 2),
        ForestColumn("participants", "N", "numeric", 1, "right"),
        ForestColumn("ci", "Effect", "ci", 3),
    )
    with pytest.raises(ValueError, match="exactly match"):
        forest(
            data,
            columns=columns,
            xlim=(0.4, 1.6),
            series_styles={"Treatment A": ForestSeriesStyle("A", "#000000")},
        )


def test_csv_rejects_nondefault_sheet_name() -> None:
    source = DATA_DIR / CASE_CSV_FILES["single_series"]
    with pytest.raises(ValueError, match="does not support sheet_name"):
        read_forest_data(source, sheet_name="Forest")


def test_csv_rejects_implicit_blank_data_row(tmp_path: Path) -> None:
    source = DATA_DIR / CASE_CSV_FILES["single_series"]
    lines = source.read_text(encoding="utf-8").splitlines()
    malformed = tmp_path / "blank-row.csv"
    malformed.write_text("\n".join([lines[0], "", *lines[1:]]) + "\n", encoding="utf-8")
    with pytest.raises(ValueError, match="entirely blank data row"):
        read_forest_data(malformed)


def test_header_legend_requires_a_column_key() -> None:
    data = load_case_data("single_series")
    columns = (
        ForestColumn("label", "Label", "text", 2),
        ForestColumn("n", "N", "numeric", 1, "right"),
        ForestColumn("effect_display", "Effect", "numeric", 2, "right"),
        ForestColumn("ci", "CI", "ci", 3),
    )
    with pytest.raises(ValueError, match="requires column_key"):
        forest(
            data,
            columns=columns,
            xlim=(-1, 1),
            legend=ForestLegendSpec(location="header"),
        )


def test_xlsx_reader_rejects_formula_without_cached_result(tmp_path: Path) -> None:
    path = tmp_path / "formula-no-cache.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Forest"
    headers = [
        "label",
        "_plot_row",
        "_series",
        "_ci_column",
        "estimate",
        "lower",
        "upper",
        "_row_type",
        "_indent",
        "_is_summary",
    ]
    sheet.append(headers)
    sheet.append(["A", "a", "S", "ci", "=1/2", 0.4, 0.6, "estimate", 0, False])
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="has no cached result"):
        read_forest_data(path)


def test_xlsx_reader_rejects_technical_field_merge(tmp_path: Path) -> None:
    path = tmp_path / "bad-merge.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Forest"
    sheet.append(
        [
            "label",
            "_plot_row",
            "_series",
            "_ci_column",
            "estimate",
            "lower",
            "upper",
            "_row_type",
            "_indent",
            "_is_summary",
        ]
    )
    sheet.append(["A", "a", "S", "ci", 0.5, 0.4, 0.6, "estimate", 0, False])
    sheet.append(["A", "b", "S", "ci", 0.6, 0.5, 0.7, "estimate", 0, False])
    sheet.merge_cells("B2:B3")
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="technical field"):
        read_forest_data(path)


def test_xlsx_reader_rejects_horizontal_display_merge(tmp_path: Path) -> None:
    path = tmp_path / "horizontal-merge.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Forest"
    sheet.append(
        [
            "label",
            "note",
            "_plot_row",
            "_series",
            "_ci_column",
            "estimate",
            "lower",
            "upper",
            "_row_type",
            "_indent",
            "_is_summary",
        ]
    )
    sheet.append(["A", "note", "a", "S", "ci", 0.5, 0.4, 0.6, "estimate", 0, False])
    sheet.merge_cells("A2:B2")
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="only single-column vertical merges"):
        read_forest_data(path)


def test_custom_mapping_is_preserved_by_reader(tmp_path: Path) -> None:
    source = DATA_DIR / CASE_CSV_FILES["single_series"]
    renamed = tmp_path / "mapped.csv"
    text = source.read_text(encoding="utf-8")
    renamed.write_text(text.replace("_plot_row", "plot_id", 1), encoding="utf-8")
    mapping = replace(ForestDataMapping(), plot_row="plot_id")
    data = read_forest_data(renamed, mapping=mapping)
    assert data.mapping.plot_row == "plot_id"
    assert "plot_id" in data.frame


def test_old_wide_api_is_removed_from_signature_and_exports() -> None:
    parameters = inspect.signature(forest).parameters
    assert {"est", "lower", "upper", "ci_column", "series_text", "series_layout"}.isdisjoint(
        parameters
    )
    assert "ForestSeriesTextSpec" not in forestploter.__all__
    assert "ForestSeriesLayoutSpec" not in forestploter.__all__
